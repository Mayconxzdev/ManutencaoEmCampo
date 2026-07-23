from __future__ import annotations

import csv
import io
import traceback
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.clock import format_local_datetime, local_today
from app.database import get_db
from app.models import Asset, MaintenanceIssue, MaintenanceOrder


router = APIRouter()


def month_competence_label(year: Optional[int], month: Optional[int]) -> Optional[str]:
    if not year or not month:
        return None
    try:
        return f"{int(month):02d}/{int(year):04d}"
    except (TypeError, ValueError):
        return None


async def generate_report_data(
    report_type: str,
    asset_id: Optional[str],
    db: AsyncSession,
    year: Optional[int] = None,
    month: Optional[int] = None,
) -> list[dict]:
    valid_types = {"schedule_monthly", "schedule_yearly", "history", "issues", "machine_individual", "assets"}
    if report_type not in valid_types:
        raise HTTPException(status_code=400, detail=f"Tipo de relatório inválido: {report_type}")

    today = local_today()
    year = year or today.year
    month = month or today.month
    if month < 1 or month > 12:
        raise HTTPException(status_code=400, detail="Mês inválido")

    if report_type == "assets":
        result = await db.execute(select(Asset).options(selectinload(Asset.sector)).order_by(Asset.type, Asset.patrimonio))
        return [{
            "Patrimônio": a.patrimonio,
            "Descrição": a.description,
            "Tipo": "Máquina" if a.type == "machine" else "Equipamento",
            "Setor": a.sector.name if a.sector else "Não informado",
            "Situação": "Ativo" if a.status == "active" else "Inativo",
            "Pode operar": "Sim" if a.is_operable else "Não",
            "Fabricante": a.manufacturer or "",
            "Modelo": a.model or "",
            "Número de série": a.serial_number or "",
            "Documentos": a.caminho_documentos or "",
            "Conflito de patrimônio": "Sim" if a.has_patrimonio_conflict else "Não",
        } for a in result.scalars().all()]

    if report_type in {"schedule_monthly", "schedule_yearly"}:
        query = select(MaintenanceOrder).options(
            selectinload(MaintenanceOrder.asset).selectinload(Asset.sector)
        ).where(MaintenanceOrder.planned_year == year)
        if report_type == "schedule_monthly":
            query = query.where(MaintenanceOrder.planned_month == month)

        result = await db.execute(query.order_by(MaintenanceOrder.planned_month, MaintenanceOrder.friendly_number))
        rows = []
        status_labels = {
            "scheduled": "Programada",
            "today": "Pendente neste mês",
            "overdue": "Atrasada",
            "in_progress": "Em andamento",
            "completed": "Concluída",
            "completed_with_issues": "Concluída com pendência",
            "cancelled": "Cancelada",
        }
        for order in result.scalars().all():
            rows.append({
                "OM": order.friendly_number,
                "Patrimônio": order.asset.patrimonio,
                "Ativo": order.asset.description,
                "Tipo": "Máquina" if order.asset.type == "machine" else "Equipamento",
                "Setor": order.asset.sector.name if order.asset.sector else "Não informado",
                "Competência": month_competence_label(order.planned_year, order.planned_month) or order.original_competence or "",
                "Data definida": order.planned_date.strftime("%d/%m/%Y") if order.scheduled_day_is_explicit and order.planned_date else "",
                "Situação": status_labels.get(order.status, order.status),
                "Responsável": order.operator_name or "",
                "Origem": order.source_status or "Aplicativo",
            })
        return rows

    if report_type in {"history", "machine_individual"}:
        query = select(MaintenanceOrder).where(
            MaintenanceOrder.status.in_(["completed", "completed_with_issues", "cancelled"])
        ).options(selectinload(MaintenanceOrder.asset))
        if report_type == "machine_individual":
            if not asset_id:
                raise HTTPException(status_code=400, detail="Selecione uma máquina para gerar este relatório.")
            query = query.where(MaintenanceOrder.asset_id == asset_id)

        result = await db.execute(query.order_by(MaintenanceOrder.completed_at.desc()))
        status_labels = {
            "completed": "Concluída",
            "completed_with_issues": "Concluída com pendência",
            "cancelled": "Cancelada",
        }
        condition_labels = {
            "released": "Liberada",
            "released_with_attention": "Liberada com atenção",
            "stopped": "Parada",
        }
        return [{
            "OM": order.friendly_number,
            "Patrimônio": order.asset.patrimonio,
            "Ativo": order.asset.description,
            "Competência": month_competence_label(order.planned_year, order.planned_month) or order.original_competence or "",
            "Conclusão": format_local_datetime(order.completed_at) if order.completed_at else "Sem data exata",
            "Conferido por": order.operator_name or "",
            "Executado por": getattr(order, "performed_by_snapshot", None) or getattr(order, "performed_by_name", None) or "",
            "Situação": status_labels.get(order.status, order.status),
            "Condição final": condition_labels.get(order.final_condition, order.final_condition or ""),
            "Observações": order.notes or order.source_note or "",
        } for order in result.scalars().all()]

    result = await db.execute(select(MaintenanceIssue).options(selectinload(MaintenanceIssue.asset)).order_by(MaintenanceIssue.created_at.desc()))
    return [{
        "Patrimônio": issue.asset.patrimonio if issue.asset else "",
        "Ativo": issue.asset.description if issue.asset else "",
        "Problema": issue.description,
        "Prioridade": issue.priority,
        "Máquina parada": "Sim" if issue.is_operable_affected else "Não",
        "Situação": "Aberta" if issue.status == "open" else "Resolvida",
        "Criada em": format_local_datetime(issue.created_at) if issue.created_at else "",
        "Resolvida em": format_local_datetime(issue.resolved_at) if issue.resolved_at else "",
        "Reparo": issue.repair_description or "",
    } for issue in result.scalars().all()]


@router.get("/reports/csv")
async def export_report_csv(
    type: str,
    asset_id: Optional[str] = None,
    year: Optional[int] = None,
    month: Optional[int] = None,
    db: AsyncSession = Depends(get_db),
):
    try:
        data = await generate_report_data(type, asset_id, db, year, month)
        stream = io.StringIO(newline="")
        if data:
            writer = csv.DictWriter(stream, fieldnames=list(data[0].keys()), delimiter=";")
            writer.writeheader()
            writer.writerows(data)
        response_bytes = ("\ufeff" + stream.getvalue()).encode("utf-8")
        response = StreamingResponse(io.BytesIO(response_bytes), media_type="text/csv; charset=utf-8")
        response.headers["Content-Disposition"] = f"attachment; filename=relatorio_{type}.csv"
        return response
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Erro ao gerar CSV: {str(exc)}") from exc


def _xlsx_from_rows(data: list[dict]) -> io.BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Relatório"
    if not data:
        sheet.append(["Nenhum registro encontrado para os filtros selecionados."])
    else:
        headers = list(data[0].keys())
        sheet.append(headers)
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in data:
            sheet.append([row.get(header, "") for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
            sheet.column_dimensions[column[0].column_letter].width = max(width, 12)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@router.get("/reports/xlsx")
async def export_report_xlsx(
    type: str,
    asset_id: Optional[str] = None,
    year: Optional[int] = None,
    month: Optional[int] = None,
    db: AsyncSession = Depends(get_db),
):
    try:
        data = await generate_report_data(type, asset_id, db, year, month)
        stream = _xlsx_from_rows(data)
        response = StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response.headers["Content-Disposition"] = f"attachment; filename=relatorio_{type}.xlsx"
        return response
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Erro ao gerar XLSX: {str(exc)}") from exc


@router.get("/reports/pdf")
async def export_report_pdf(
    type: str,
    asset_id: Optional[str] = None,
    order_id: Optional[str] = None,
    year: Optional[int] = None,
    month: Optional[int] = None,
    db: AsyncSession = Depends(get_db),
):
    try:
        from app.pdf_generator import (
            generate_asset_history_pdf,
            generate_complete_order_pdf,
            generate_issues_pdf,
            generate_monthly_schedule_pdf,
            generate_yearly_schedule_pdf,
        )

        if type == "schedule_monthly":
            pdf_buf = await generate_monthly_schedule_pdf(db, year, month)
            filename = "programacao_mensal.pdf"
        elif type == "schedule_yearly":
            pdf_buf = await generate_yearly_schedule_pdf(db, year)
            filename = "programacao_anual.pdf"
        elif type == "history":
            pdf_buf = await generate_asset_history_pdf(db, asset_id)
            filename = f"historico_{asset_id or 'geral'}.pdf"
        elif type == "issues":
            pdf_buf = await generate_issues_pdf(db, stopped_only=False)
            filename = "pendencias_ocorrencias.pdf"
        elif type == "stopped_machines":
            pdf_buf = await generate_issues_pdf(db, stopped_only=True)
            filename = "maquinas_paradas.pdf"
        elif type == "order_complete":
            if not order_id:
                raise HTTPException(status_code=400, detail="Parâmetro order_id é obrigatório para relatórios do tipo order_complete")
            pdf_buf = await generate_complete_order_pdf(db, order_id)
            filename = f"ordem_manutencao_{order_id}.pdf"
        else:
            raise HTTPException(status_code=400, detail=f"Tipo de relatório PDF inválido: {type}")

        response = StreamingResponse(pdf_buf, media_type="application/pdf")
        response.headers["Content-Disposition"] = f"attachment; filename={filename}"
        return response
    except HTTPException:
        raise
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except Exception as exc:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Erro ao gerar PDF: {str(exc)}") from exc


@router.get("/reports/orders/{order_id}/form.pdf")
async def get_order_form_pdf(order_id: str, db: AsyncSession = Depends(get_db)):
    try:
        from app.pdf_generator import generate_complete_order_pdf

        pdf_buf = await generate_complete_order_pdf(db, order_id)
        response = StreamingResponse(pdf_buf, media_type="application/pdf")
        response.headers["Content-Disposition"] = f"inline; filename=ordem_{order_id}.pdf"
        return response
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc)) from exc
